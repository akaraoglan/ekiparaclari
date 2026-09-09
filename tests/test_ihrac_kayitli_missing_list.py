import io
import os
import unittest
import uuid
from contextlib import contextmanager
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

import app as app_module
from tools.ihrac_kayitli import create_missing_invoice_excel


@contextmanager
def _flat_workspace(root):
    root.mkdir(parents=True, exist_ok=True)
    existing_files = {path for path in root.iterdir() if path.is_file()}
    try:
        yield str(root)
    finally:
        for path in root.iterdir():
            if path.is_file() and path not in existing_files:
                path.unlink(missing_ok=True)


class IhracKayitliMissingListTest(unittest.TestCase):
    TEST_ROOT = Path(
        os.environ.get(
            "EKIPARACLARI_TEST_TMPDIR",
            Path(__file__).resolve().parents[1] / "outputs",
        )
    )

    def test_excel_has_no_header_or_xml_suffix(self):
        references = ["SE02026000003655.xml", "SE12026000002647"]
        with _flat_workspace(self.TEST_ROOT) as folder:
            output_path = create_missing_invoice_excel(references, folder)
            workbook = load_workbook(output_path)
            worksheet = workbook["Eksik Faturalar"]

            self.assertEqual("SE02026000003655", worksheet["A1"].value)
            self.assertEqual("SE12026000002647", worksheet["A2"].value)
            self.assertEqual(2, worksheet.max_row)
            self.assertEqual(1, worksheet.max_column)
            self.assertEqual("@", worksheet["A1"].number_format)
            workbook.close()

    def test_web_route_downloads_still_missing_invoices(self):
        batch_id = uuid.uuid4().hex
        with _flat_workspace(self.TEST_ROOT) as folder:
            self._write_summary(Path(folder) / f"{batch_id}_ozet.xlsx")
            self._write_detail(Path(folder) / f"{batch_id}_detay.xlsx")

            app_module.app.config.update(TESTING=True)
            with patch.object(app_module, "UPLOAD_DIR", folder), patch.object(
                app_module,
                "OUTPUT_DIR",
                folder,
            ):
                response = app_module.app.test_client().post(
                    "/starwood/ihrac-kayitli-eksik-xml-listesi",
                    data={"batch_id": batch_id},
                )

            self.assertEqual(200, response.status_code)
            self.assertIn("attachment", response.headers["Content-Disposition"])
            self.assertIn("eksik_fatura_listesi.xlsx", response.headers["Content-Disposition"])
            workbook = load_workbook(io.BytesIO(response.data))
            worksheet = workbook["Eksik Faturalar"]
            self.assertEqual("SE02026000003655", worksheet["A1"].value)
            self.assertEqual("SE12026000002647", worksheet["A2"].value)
            self.assertEqual(2, worksheet.max_row)
            self.assertEqual(1, worksheet.max_column)
            workbook.close()
            response.close()

    @staticmethod
    def _write_summary(path):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["Faturalama tarihi", "Referans", "Fatura alıcısı adı", "Vergi Numarası"])
        worksheet.append(["2026-09-09", "SE12026000002647", "Alıcı 2", "222"])
        worksheet.append(["2026-09-09", "SE02026000003655", "Alıcı 1", "111"])
        workbook.save(path)
        workbook.close()

    @staticmethod
    def _write_detail(path):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append([
            "Referans",
            "KDV Matrahı",
            "Belge para birimi",
            "Faturalanan miktar",
            "Kalınlık",
            "Boy",
            "En",
            "Gtip",
        ])
        worksheet.append(["SE12026000002647", 100, "USD", 1, 1, 1, 1, "A"])
        worksheet.append(["SE02026000003655", 200, "EUR", 1, 1, 1, 1, "B"])
        workbook.save(path)
        workbook.close()


if __name__ == "__main__":
    unittest.main()
