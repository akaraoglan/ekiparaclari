import unittest
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch
from types import SimpleNamespace

import fitz
from openpyxl import load_workbook

from tools.isci_vergi_toplami import (
    TextBox,
    OCRUnavailableError,
    _native_boxes,
    _ocr_banded_boxes,
    _EXCEL_COLUMNS,
    _extract_employee_count,
    _canonical_extra_payment_label,
    _extract_extra_payment_items,
    _find_extra_payment_total,
    _named_amount,
    analyze_tax_pdfs,
    create_tax_excel,
    extract_page_values,
)


class IsciVergiToplamiTest(unittest.TestCase):
    def test_reads_employee_count_without_using_neighboring_counts(self):
        lines = [[
            TextBox("ÇALIŞAN SAYISI", 10, 60, 130, 70),
            TextBox(":", 140, 60, 145, 70),
            TextBox("21", 180, 60, 200, 70),
            TextBox("İşe Giren : 7", 310, 60, 400, 70),
        ]]
        self.assertEqual(21, _extract_employee_count(lines, 1000, 1500))
        self.assertEqual(1109, _extract_employee_count(
            [[TextBox("ÇALIŞAN SAYISI : 1.109", 10, 60, 250, 70)]], 1000, 1500
        ))
        self.assertEqual(0, _extract_employee_count(
            [[TextBox("CALISAN SAYISI : 0", 10, 60, 200, 70)]], 1000, 1500
        ))
        self.assertIsNone(_extract_employee_count(
            [[TextBox("Kadın Sayısı : 8", 10, 60, 200, 70)]], 1000, 1500
        ))

    def test_extracts_each_extra_payment_and_reconciles_total(self):
        lines = [
            self._line(200, "Çocuk Parası", "3.544,92"),
            self._line(240, "Yıllık İzin Harçlığı", "5.730,00"),
            self._line(280, "Yakacak Yardımı (Na", "368.678,04"),
        ]

        payments, metadata = _extract_extra_payment_items(
            lines,
            width=1000,
            extra_y=150,
            total_y=350,
        )

        self.assertEqual(Decimal("3544.92"), payments["ocr_çocuk parası"])
        self.assertEqual(Decimal("5730.00"), payments["ocr_yıllık izin harçlığı"])
        self.assertEqual(Decimal("368678.04"), payments["ocr_yakacak yardımı (na"])
        self.assertEqual(Decimal("377952.96"), sum(payments.values()))
        self.assertFalse(metadata["ocr_yakacak yardımı (na"]["estimated"])

    def test_preserves_payment_headings_without_guessing_categories(self):
        for label in (
            "Bayram Yardımı", "Yemek Yardımı", "Yemek Ödemesi",
            "a ik Yardimi (Na", "Yakacak Yardımı (Na",
            "Doğum Yardımı", "Doğum Yardımı (Nakdi)", "Do um Yardimi",
            "� Yardımı",
        ):
            with self.subTest(label=label):
                _, header, _, _, _ = _canonical_extra_payment_label(label)
                self.assertEqual(label, header)
        _, header, needs_review, reason, _ = _canonical_extra_payment_label(
            "Bayram Yardımı", confidence=0.4,
        )
        self.assertEqual("Bayram Yardımı", header)
        self.assertTrue(needs_review)
        self.assertIn("tahmin edilmedi", reason)

    def test_missing_heading_keeps_amount_with_review_warning(self):
        payments, metadata = _extract_extra_payment_items(
            [self._line(200, "", "125,00")], 1000, 150, 250,
        )
        self.assertEqual(Decimal("125.00"), sum(payments.values()))
        self.assertTrue(next(iter(metadata.values()))["estimated"])
        self.assertIn("Başlık okunamadı", next(iter(metadata.values()))["header"])

    def test_similar_headings_stay_separate_across_pdfs_and_excel(self):
        output_root = Path(__file__).resolve().parents[1] / "outputs"
        output_root.mkdir(parents=True, exist_ok=True)
        pdf_path = output_root / "test_extra_payments.pdf"
        with fitz.open() as document:
            document.new_page()
            document.new_page()
            document.save(pdf_path)
        values = []
        expected = {
            "Bayram Yardımı": Decimal("100.00"),
            "Yemek Yardımı": Decimal("200.00"),
            "Doğum Yardımı": Decimal("300.00"),
            "Do um Yardimi": Decimal("400.00"),
            "Yemek Yardımı (Nakdi)": Decimal("500.00"),
        }
        for labels in (list(expected)[:3], list(expected)[3:]):
            payments, metadata = _extract_extra_payment_items(
                [self._line(200 + i * 20, label, str(expected[label]))
                 for i, label in enumerate(labels)], 1000, 150, 350,
            )
            values.append({
                **{key: Decimal("0") for key, _ in _EXCEL_COLUMNS},
                "employee_count": 1,
                "extra_payments": payments, "extra_payment_meta": metadata,
            })
        excel_path = None
        try:
            with patch("tools.isci_vergi_toplami.extract_page_values", side_effect=values):
                result = analyze_tax_pdfs([(str(pdf_path), "ornek.pdf")])
            self.assertEqual(list(expected), [c["header"] for c in result["extra_payment_columns"]])
            self.assertEqual(sum(expected.values()), sum(result["extra_payment_totals"].values()))
            self.assertEqual([], result["warnings"])
            excel_path = create_tax_excel(result, str(output_root))
            workbook = load_workbook(excel_path)
            try:
                sheet = workbook["Vergi Toplamları"]
                start = 4 + len(_EXCEL_COLUMNS)
                for offset, (label, amount) in enumerate(expected.items()):
                    column = start + offset
                    self.assertEqual(label, sheet.cell(3, column).value)
                    row = 4 if offset < 3 else 5
                    self.assertEqual(float(amount), sheet.cell(row, column).value)
                    self.assertEqual(0, sheet.cell(9 - row, column).value)
            finally:
                workbook.close()
        finally:
            pdf_path.unlink(missing_ok=True)
            if excel_path:
                Path(excel_path).unlink(missing_ok=True)

    def test_finds_total_when_first_letter_is_not_read(self):
        lines = [self._line(400, "oplam", "377.952,96")]

        result = _find_extra_payment_total(
            lines,
            width=1000,
            y_min=300,
            y_max=500,
            target_x=400,
        )

        self.assertEqual((405.0, Decimal("377952.96")), result)

    def test_reads_wrapped_amount_from_line_below_heading(self):
        lines = [
            [TextBox("İşveren Maliyeti (Teşvikli)", 450, 700, 760, 710)],
            [TextBox("879.691,50", 830, 716, 950, 726)],
        ]

        amount = _named_amount(
            lines,
            "isveren maliyeti",
            width=1000,
            label_x=(0.43, 0.78),
            amount_x=(0.78, 0.99),
            y_min=600,
            y_max=900,
            allow_next_line=True,
        )

        self.assertEqual(Decimal("879691.50"), amount)

    def test_scanned_page_retries_at_higher_resolution(self):
        expected = {"gross": Decimal("1.00")}
        with patch(
            "tools.isci_vergi_toplami._native_boxes",
            return_value=[],
        ), patch(
            "tools.isci_vergi_toplami._ocr_boxes",
            side_effect=[([], 1000.0, 1500.0), ([], 1500.0, 2250.0)],
        ) as ocr_boxes, patch(
            "tools.isci_vergi_toplami._extract_page_values_from_boxes",
            side_effect=[ValueError("düşük kalite"), expected],
        ):
            result = extract_page_values(object())

        self.assertIs(expected, result)
        self.assertEqual(
            [2.0, 3.0],
            [call.kwargs["scale"] for call in ocr_boxes.call_args_list],
        )

    def test_native_pdf_with_joined_section_title_reads_second_page_without_ocr(self):
        output_root = Path(__file__).resolve().parents[1] / "outputs"
        output_root.mkdir(parents=True, exist_ok=True)
        pdf_path = output_root / "test_native_payroll.pdf"
        excel_path = None
        try:
            with fitz.open() as doc:
                self._add_payroll_page(doc, "Ek Odemeler")
                self._add_payroll_page(doc, "EkOdemeler")
                doc.save(pdf_path)
            with patch("tools.isci_vergi_toplami._ocr_boxes") as ocr:
                result = analyze_tax_pdfs([(str(pdf_path), "ornek.pdf")])
                ocr.assert_not_called()
            self.assertEqual([1, 2], [row["page"] for row in result["rows"]])
            self.assertEqual([], result["failed_pages"])
            self.assertEqual([], result["warnings"])
            self.assertEqual(Decimal("2000.00"), result["gross_total"])
            self.assertEqual(Decimal("200.00"), result["income_total"])
            excel_path = create_tax_excel(result, str(output_root))
            workbook = load_workbook(excel_path)
            try:
                sheet = workbook["Vergi Toplamları"]
                self.assertEqual(2, sheet["B5"].value)
                self.assertEqual(1000, sheet["D5"].value)
                self.assertEqual("=SUM(D4:D5)", sheet["D6"].value)
            finally:
                workbook.close()
        finally:
            pdf_path.unlink(missing_ok=True)
            if excel_path:
                Path(excel_path).unlink(missing_ok=True)

    def test_banded_ocr_recovers_after_full_page_reads_fail(self):
        page = SimpleNamespace(rect=fitz.Rect(0, 0, 1000, 1500))
        expected = {"gross": Decimal("123.45")}
        with patch("tools.isci_vergi_toplami._native_boxes", return_value=[]), patch(
            "tools.isci_vergi_toplami._ocr_boxes", return_value=([], 2000, 3000),
        ), patch(
            "tools.isci_vergi_toplami._ocr_banded_boxes", return_value=([], 3000, 4500),
        ) as banded, patch(
            "tools.isci_vergi_toplami._extract_page_values_from_boxes",
            side_effect=[ValueError("native"), ValueError("2x"), ValueError("3x"), expected],
        ):
            self.assertIs(expected, extract_page_values(page))
            banded.assert_called_once_with(page)

    def test_bands_keep_overlap_rows_once_in_page_coordinates(self):
        page = SimpleNamespace(rect=fitz.Rect(0, 0, 1000, 1000))
        boxes = [TextBox(str(y), 10, y - 5, 30, y + 5) for y in (249, 250, 251, 500, 750)]
        clips = []
        def read_region(page, clip, scale):
            clips.append(clip)
            return [box for box in boxes if clip.y0 <= box.center_y <= clip.y1]
        with patch("tools.isci_vergi_toplami._ocr_region_boxes", side_effect=read_region):
            actual, width, height = _ocr_banded_boxes(page, scale=1.0)
        self.assertEqual([box.text for box in boxes], [box.text for box in actual])
        self.assertEqual((1000, 1000), (width, height))
        self.assertGreater(clips[0].y1, clips[1].y0)

    def test_page_inference_error_retries_but_missing_ocr_dependencies_abort(self):
        page = SimpleNamespace(rect=fitz.Rect(0, 0, 1000, 1500))
        expected = {"gross": Decimal("123.45")}
        with patch("tools.isci_vergi_toplami._native_boxes", return_value=[]), patch(
            "tools.isci_vergi_toplami._ocr_boxes",
            side_effect=[RuntimeError("image failure"), ([], 3000, 4500)],
        ), patch(
            "tools.isci_vergi_toplami._extract_page_values_from_boxes",
            side_effect=[ValueError("native"), expected],
        ):
            self.assertIs(expected, extract_page_values(page))
        with patch("tools.isci_vergi_toplami._native_boxes", return_value=[]), patch(
            "tools.isci_vergi_toplami._ocr_boxes", side_effect=OCRUnavailableError("missing model"),
        ) as ocr:
            with self.assertRaisesRegex(OCRUnavailableError, "missing model"):
                extract_page_values(page)
            self.assertEqual(1, ocr.call_count)

    def test_all_reading_failures_are_reported_with_strategy_names(self):
        page = SimpleNamespace(rect=fitz.Rect(0, 0, 1000, 1500))
        with patch("tools.isci_vergi_toplami._native_boxes", return_value=[]), patch(
            "tools.isci_vergi_toplami._ocr_boxes", return_value=([], 2000, 3000),
        ), patch(
            "tools.isci_vergi_toplami._ocr_banded_boxes", return_value=([], 3000, 4500),
        ):
            with self.assertRaises(ValueError) as failure:
                extract_page_values(page)
        for name in ("PDF metni", "Tam sayfa OCR (2x)", "Tam sayfa OCR (3x)", "Bölgesel OCR"):
            self.assertIn(name, str(failure.exception))

    def test_failed_second_page_is_visible_and_third_page_still_contributes(self):
        output_root = Path(__file__).resolve().parents[1] / "outputs"
        output_root.mkdir(parents=True, exist_ok=True)
        pdf_path = output_root / "test_failed_payroll.pdf"
        excel_path = None
        try:
            with fitz.open() as doc:
                for _ in range(3):
                    doc.new_page()
                doc.save(pdf_path)
            values = {key: Decimal("10.00") for key, _ in _EXCEL_COLUMNS}
            values["employee_count"] = 1
            with patch("tools.isci_vergi_toplami.extract_page_values", side_effect=[
                dict(values), ValueError("Ek Ödemeler toplamı okunamadı"), dict(values),
            ]):
                result = analyze_tax_pdfs([(str(pdf_path), "ornek.pdf")])
            self.assertEqual([1, 3], [row["page"] for row in result["rows"]])
            self.assertEqual(2, result["failed_pages"][0]["page"])
            self.assertEqual(Decimal("20.00"), result["gross_total"])
            self.assertIn("toplamlar eksiktir", result["warnings"][0])
            excel_path = create_tax_excel(result, str(output_root))
            workbook = load_workbook(excel_path)
            try:
                sheet = workbook["Vergi Toplamları"]
                self.assertIn("1 sayfa okunamadı", sheet["A2"].value)
                self.assertEqual("OKUNAN SAYFALAR TOPLAMI", sheet["A6"].value)
                self.assertIn("Sayfa 2", workbook["Uyarılar"]["A3"].value)
            finally:
                workbook.close()
        finally:
            pdf_path.unlink(missing_ok=True)
            if excel_path:
                Path(excel_path).unlink(missing_ok=True)

    def test_rotated_native_payroll_reads_without_changing_original_page(self):
        with fitz.open() as doc:
            page = self._add_payroll_page(doc, "EkOdemeler")
            page.set_rotation(90)
            with patch("tools.isci_vergi_toplami._ocr_boxes") as ocr:
                result = extract_page_values(page)
                ocr.assert_not_called()
            self.assertEqual(Decimal("1000.00"), result["gross"])
            self.assertEqual(Decimal("100.00"), result["income"])
            self.assertEqual(90, page.rotation)

    def test_native_coordinates_match_rotated_page_rendering(self):
        with fitz.open() as doc:
            page = doc.new_page(width=600, height=900)
            page.insert_text((50, 100), "Test", fontsize=10)
            rect = fitz.Rect(page.get_text("words")[0][:4])
            page.set_rotation(90)
            box = _native_boxes(page)[0]
            self.assertEqual(tuple(rect * page.rotation_matrix), (box.x0, box.y0, box.x1, box.y1))

    @staticmethod
    def _add_payroll_page(doc, extra_heading):
        page = doc.new_page(width=1000, height=1500)
        def put(x, y, text):
            page.insert_text((x, y), text, fontsize=10)
        def left(y, text, value=None):
            put(10, y, text)
            if value is not None:
                put(370, y, value)
        def right(y, text, worker=None, employer=None):
            put(450, y, text)
            if worker is not None:
                put(790, y, worker)
            if employer is not None:
                put(910, y, employer)
        left(60, "CALISAN SAYISI: 21")
        left(100, "Calismalar")
        put(370, 100, "Brut Tutar")
        put(790, 100, "Isci")
        put(910, 100, "Isveren")
        for args in [
            (300, "Toplam", "1.000,00"), (350, "Fazla Mesailer"),
            (400, "Toplam", "50,00"), (450, extra_heading),
            (500, "Bayram Yardimi", "100,00"), (540, "Yemek Yardimi", "200,00"),
            (580, "Toplam", "300,00"), (620, "Brut Odemeler"),
        ]:
            left(*args)
        for args in [
            (150, "SGK", "140,00", "205,00"), (180, "SGDP", "0,00", "0,00"),
            (210, "Issizlik", "10,00", "20,00"), (240, "Ek Issizlik", "0,00", "0,00"),
            (270, "Gelir Vergisi", "100,00"), (300, "Damga Vergisi", "10,00"),
            (330, "Yasal Kesintiler Toplami", None, "225,00"), (370, "Ozel Kesintiler"),
            (400, "Avans", None, "50,00"), (430, "Toplam", None, "50,00"),
            (470, "Vergi Indirimi"), (500, "Gelir Vergisi Indirimi", None, "20,00"),
            (530, "Damga Vergisi Indirimi", None, "2,00"), (570, "Tesvikler"),
            (600, "Tesvik A", None, "15,00"),
            (650, "Isveren Maliyeti (Tesvikli)", None, "1.560,00"),
            (700, "Net Odenen", None, "1.040,00"), (730, "Otomatik BES", None, "5,00"),
        ]:
            right(*args)
        return page

    def test_excel_appends_dynamic_columns_and_keeps_amounts_numeric(self):
        rows = []
        for page in (1, 2):
            row = {
                "filename": "ornek.pdf",
                "page": page,
                "employee_count": 21 if page == 1 else None,
                **{key: Decimal("1.00") for key, _ in _EXCEL_COLUMNS},
                "extra_payments": {
                    "cocuk_parasi": Decimal("10.00") * page,
                    "yakacak_yardimi": Decimal("20.00") * page,
                },
            }
            rows.append(row)

        result = {
            "rows": rows,
            "warnings": ["Yakacak yardımı başlığı kontrol edilmelidir."],
            "extra_payment_columns": [
                {
                    "key": "cocuk_parasi",
                    "header": "Çocuk Parası",
                    "estimated_sources": [],
                },
                {
                    "key": "yakacak_yardimi",
                    "header": "Yakacak Yardımı",
                    "estimated_sources": [{"row_index": 1}],
                },
            ],
        }

        output_root = Path(__file__).resolve().parents[1] / "outputs"
        output_root.mkdir(parents=True, exist_ok=True)
        output_path = create_tax_excel(result, str(output_root))
        try:
            workbook = load_workbook(output_path, data_only=False)
            try:
                sheet = workbook["Vergi Toplamları"]
                self.assertEqual("Çalışan Sayısı", sheet["C3"].value)
                self.assertEqual(21, sheet["C4"].value)
                self.assertEqual("n", sheet["C4"].data_type)
                self.assertEqual("#,##0", sheet["C4"].number_format)
                self.assertIsNone(sheet["C5"].value)
                self.assertEqual("=SUM(C4:C5)", sheet["C6"].value)
                self.assertEqual("#,##0", sheet["C6"].number_format)
                self.assertEqual("Brüt Ücret (YK Dahil)", sheet["D3"].value)
                self.assertEqual("İşveren Maliyeti (Teşvikli)", sheet["U3"].value)
                self.assertEqual("Çocuk Parası", sheet["V3"].value)
                self.assertIn("Kontrol gerekli - Excel satırı: 5", sheet["W3"].value)
                self.assertEqual("F4CCCC", sheet["W3"].fill.fgColor.rgb[-6:])
                self.assertEqual(40.0, sheet["W5"].value)
                self.assertEqual("n", sheet["W5"].data_type)
                self.assertEqual("#,##0.00", sheet["W5"].number_format)
                self.assertEqual("=SUM(W4:W5)", sheet["W6"].value)
                self.assertIn("Uyarılar", workbook.sheetnames)
            finally:
                workbook.close()
        finally:
            Path(output_path).unlink(missing_ok=True)

    @staticmethod
    def _line(y: float, label: str, amount: str) -> list[TextBox]:
        return [
            TextBox(label, 10, y, 260, y + 10),
            TextBox(amount, 350, y, 450, y + 10),
        ]


if __name__ == "__main__":
    unittest.main()
