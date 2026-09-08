import os
import re
import threading
import unicodedata
import uuid
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from statistics import median

import fitz
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


_OCR_ENGINE = None
_OCR_LOCK = threading.Lock()
_MONEY_RE = re.compile(r"(?<!\d)\d{1,3}(?:\.\d{3})*,\d{2}(?!\d)|(?<!\d)\d+,\d{2}(?!\d)")
_EXCEL_COLUMNS = [
    ("gross", "Brüt Ücret (YK Dahil)"),
    ("income", "Gelir Vergisi (İşçi)"),
    ("stamp", "Damga Vergisi (İşçi)"),
    ("overtime_total", "Fazla Mesailer Toplamı"),
    ("bonus", "İkramiye"),
    ("other_extra_payments", "Diğer Ek Ödeme ve Yardımlar"),
    ("worker_sgk", "İşçi SGK"),
    ("worker_unemployment", "İşçi İşsizlik"),
    ("employer_sgk", "İşveren SGK"),
    ("employer_unemployment", "İşveren İşsizlik"),
    ("income_tax_incentive", "Gelir Vergisi Teşviki"),
    ("stamp_tax_incentive", "Damga Vergisi Teşviki"),
    ("advance", "Avans"),
    ("other_advances", "Avans 1-2-3"),
    ("account_369_deductions", "369 Kesintiler"),
    ("net_paid", "Net Ödenen"),
    ("incentives_total", "Teşvikler Toplamı"),
    ("employer_cost", "İşveren Maliyeti (Teşvikli)"),
]


@dataclass(frozen=True)
class TextBox:
    text: str
    x0: float
    y0: float
    x1: float
    y1: float

    @property
    def center_x(self) -> float:
        return (self.x0 + self.x1) / 2

    @property
    def center_y(self) -> float:
        return (self.y0 + self.y1) / 2


def _plain(text: str) -> str:
    text = text.casefold().replace("ı", "i")
    text = "".join(
        char for char in unicodedata.normalize("NFKD", text)
        if not unicodedata.combining(char)
    )
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def _parse_money(text: str) -> Decimal | None:
    compact = text.replace(" ", "")
    match = _MONEY_RE.search(compact)
    if not match:
        zero_candidate = compact.upper().replace("O", "0")
        if re.fullmatch(r"0+[,.']0+", zero_candidate):
            return Decimal("0")
        return None
    try:
        return Decimal(match.group(0).replace(".", "").replace(",", "."))
    except InvalidOperation:
        return None


def format_tr_money(value: Decimal) -> str:
    rendered = f"{value:,.2f}"
    return rendered.replace(",", "_").replace(".", ",").replace("_", ".")


def _native_boxes(page) -> list[TextBox]:
    return [
        TextBox(str(word[4]), float(word[0]), float(word[1]), float(word[2]), float(word[3]))
        for word in page.get_text("words")
        if str(word[4]).strip()
    ]


def _run_rapidocr(image):
    try:
        from rapidocr import RapidOCR
    except ImportError as exc:
        raise RuntimeError(
            "OCR paketleri eksik. Sunucuda 'pip install -r requirements.txt' komutunu çalıştırın."
        ) from exc

    global _OCR_ENGINE
    try:
        with _OCR_LOCK:
            if _OCR_ENGINE is None:
                _OCR_ENGINE = RapidOCR()
            result = _OCR_ENGINE(image)
    except Exception as exc:
        detail = str(exc).splitlines()[0][:160]
        if "libgl" in detail.casefold():
            raise RuntimeError(
                "OCR başlatılamadı. Linux sunucuda libGL1 paketini kurup servisi yeniden başlatın."
            ) from exc
        raise RuntimeError(f"OCR çalıştırılamadı ({type(exc).__name__}: {detail}).") from exc
    return result


def _ocr_region_boxes(page, clip: fitz.Rect, scale: float) -> list[TextBox]:
    try:
        import numpy as np
    except ImportError as exc:
        raise RuntimeError(
            "OCR paketleri eksik. Sunucuda 'pip install -r requirements.txt' komutunu çalıştırın."
        ) from exc

    pixmap = page.get_pixmap(
        matrix=fitz.Matrix(scale, scale),
        clip=clip,
        colorspace=fitz.csRGB,
        alpha=False,
    )
    image = np.frombuffer(pixmap.samples, dtype=np.uint8).reshape(
        pixmap.height, pixmap.width, pixmap.n
    )
    result = _run_rapidocr(image)
    offset_x = clip.x0 * scale
    offset_y = clip.y0 * scale

    boxes = []
    if result.boxes is not None:
        for points, text in zip(result.boxes, result.txts):
            xs = [float(point[0]) + offset_x for point in points]
            ys = [float(point[1]) + offset_y for point in points]
            boxes.append(TextBox(str(text), min(xs), min(ys), max(xs), max(ys)))
    return boxes


def _ocr_boxes(page, scale: float = 2.0) -> tuple[list[TextBox], float, float]:
    width = float(page.rect.width * scale)
    height = float(page.rect.height * scale)
    boxes = _ocr_region_boxes(page, page.rect, scale)
    return boxes, width, height


def _group_lines(boxes: list[TextBox], page_height: float) -> list[list[TextBox]]:
    lines: list[list[TextBox]] = []
    box_heights = [box.y1 - box.y0 for box in boxes]
    typical_height = median(box_heights) if box_heights else 8.0
    tolerance = max(3.0, typical_height * 0.48)
    for box in sorted(boxes, key=lambda item: (item.center_y, item.x0)):
        best_line = None
        best_distance = None
        for line in lines:
            line_y = sum(item.center_y for item in line) / len(line)
            distance = abs(box.center_y - line_y)
            if distance <= tolerance and (best_distance is None or distance < best_distance):
                best_line = line
                best_distance = distance
        if best_line is None:
            lines.append([box])
        else:
            best_line.append(box)
    for line in lines:
        line.sort(key=lambda item: item.x0)
    return sorted(lines, key=lambda line: sum(item.center_y for item in line) / len(line))


def _line_text(line: list[TextBox]) -> str:
    return " ".join(item.text for item in line)


def _line_y(line: list[TextBox]) -> float:
    return sum(item.center_y for item in line) / len(line)


def _regional_text(
    line: list[TextBox], width: float, x_min: float, x_max: float
) -> str:
    return " ".join(
        item.text for item in line
        if width * x_min <= item.center_x <= width * x_max
    )


def _regional_amounts(
    line: list[TextBox], width: float, x_min: float, x_max: float
) -> list[tuple[float, Decimal]]:
    return [
        (item.center_x, amount)
        for item in line
        if width * x_min <= item.center_x <= width * x_max
        and (amount := _parse_money(item.text)) is not None
    ]


def _find_section_y(
    lines: list[list[TextBox]],
    label: str,
    width: float,
    x_min: float,
    x_max: float,
    after_y: float = 0,
) -> float | None:
    wanted = _plain(label)
    for line in lines:
        if _line_y(line) <= after_y:
            continue
        if wanted in _plain(_regional_text(line, width, x_min, x_max)):
            return _line_y(line)
    return None


def _named_amount(
    lines: list[list[TextBox]],
    label: str,
    width: float,
    label_x: tuple[float, float],
    amount_x: tuple[float, float],
    y_min: float = 0,
    y_max: float = float("inf"),
) -> Decimal | None:
    wanted = _plain(label)
    for line in lines:
        line_y = _line_y(line)
        if not y_min < line_y < y_max:
            continue
        line_label = _plain(_regional_text(line, width, *label_x))
        if wanted not in line_label:
            continue
        amounts = _regional_amounts(line, width, *amount_x)
        if amounts:
            return max(amounts, key=lambda pair: pair[0])[1]
    return None


def _section_total(
    lines: list[list[TextBox]],
    width: float,
    y_min: float,
    y_max: float,
    label_x: tuple[float, float],
    amount_x: tuple[float, float],
    target_x: float | None = None,
) -> Decimal | None:
    for line in lines:
        line_y = _line_y(line)
        if not y_min < line_y < y_max:
            continue
        if "toplam" not in _plain(_regional_text(line, width, *label_x)):
            continue
        amounts = _regional_amounts(line, width, *amount_x)
        if not amounts:
            continue
        if target_x is None:
            return max(amounts, key=lambda pair: pair[0])[1]
        return min(amounts, key=lambda pair: abs(pair[0] - target_x))[1]
    return None


def _find_worker_x(lines: list[list[TextBox]], width: float) -> float:
    for line in lines:
        for item in line:
            normalized = _plain(item.text).replace("1", "i")
            if normalized in {"isci", "isci i"} or "isci" in normalized:
                if item.center_y < max(box.center_y for row in lines for box in row) * 0.55:
                    return item.center_x
    return width * 0.80


def _find_employer_x(lines: list[list[TextBox]], width: float) -> float:
    for line in lines:
        for item in line:
            if "isveren" in _plain(item.text).replace("1", "i"):
                return item.center_x
    return width * 0.92


def _deduction_row_amount(
    lines: list[list[TextBox]],
    label: str,
    target_x: float,
    width: float,
    height: float,
) -> Decimal | None:
    wanted = _plain(label)
    for line in lines:
        if _line_y(line) >= height * 0.30:
            break
        label_items = [
            item for item in line
            if width * 0.43 <= item.center_x <= width * 0.61
        ]
        if not any(_plain(item.text) == wanted for item in label_items):
            continue
        amounts = _regional_amounts(line, width, 0.68, 0.99)
        if amounts:
            return min(amounts, key=lambda pair: abs(pair[0] - target_x))[1]
    return None


def _amount_on_named_line(
    lines: list[list[TextBox]], label: str, worker_x: float, width: float
) -> Decimal | None:
    label_plain = _plain(label)
    for line in lines:
        text_plain = _plain(_line_text(line))
        if label_plain not in text_plain:
            continue
        candidates = [
            (abs(item.center_x - worker_x), amount)
            for item in line
            if (amount := _parse_money(item.text)) is not None
            and width * 0.69 <= item.center_x <= width * 0.88
        ]
        if candidates:
            return min(candidates, key=lambda pair: pair[0])[1]
    return None


def _fallback_worker_amounts(
    lines: list[list[TextBox]], worker_x: float, width: float
) -> tuple[Decimal | None, Decimal | None]:
    header_y = None
    for line in lines:
        if any("isci" in _plain(item.text).replace("1", "i") for item in line):
            header_y = sum(item.center_y for item in line) / len(line)
            break
    if header_y is None:
        return None, None

    amounts: list[tuple[float, Decimal]] = []
    for line in lines:
        line_y = sum(item.center_y for item in line) / len(line)
        if line_y <= header_y:
            continue
        candidates = [
            (abs(item.center_x - worker_x), _parse_money(item.text))
            for item in line
            if width * 0.69 <= item.center_x <= width * 0.88
            and _parse_money(item.text) is not None
        ]
        if candidates:
            amounts.append((line_y, min(candidates, key=lambda pair: pair[0])[1]))
        if len(amounts) >= 5:
            break
    if len(amounts) >= 5:
        return amounts[3][1], amounts[4][1]
    return None, None


def _extract_gross_amount(
    lines: list[list[TextBox]], width: float, height: float
) -> tuple[Decimal, Decimal]:
    gross_header_x = width * 0.40
    gross_header_y = None
    for line in lines:
        line_text = _plain(_line_text(line))
        if _line_y(line) < height * 0.25 and "brut tutar" in line_text:
            gross_header_y = sum(item.center_y for item in line) / len(line)
            matching = [item for item in line if "brut" in _plain(item.text)]
            if matching:
                gross_header_x = matching[0].center_x
            break

    if gross_header_y is None:
        gross_header_y = _find_section_y(
            lines, "calismalar", width, 0.00, 0.52
        )
    if gross_header_y is None:
        raise ValueError("Çalışmalar bölümünün başlığı okunamadı.")

    base_gross = None
    for line in lines:
        line_y = sum(item.center_y for item in line) / len(line)
        if line_y <= gross_header_y or line_y >= height * 0.36:
            continue
        if "toplam" not in _plain(_line_text(line)):
            continue
        candidates = [
            (abs(item.center_x - gross_header_x), amount)
            for item in line
            if (amount := _parse_money(item.text)) is not None
            and width * 0.34 <= item.center_x <= width * 0.50
        ]
        if candidates:
            base_gross = min(candidates, key=lambda pair: pair[0])[1]
            break

    if base_gross is None:
        raise ValueError("Çalışmalar bölümündeki toplam brüt tutar okunamadı.")

    yk_fee = Decimal("0")
    for line in lines:
        normalized = _plain(_line_text(line))
        if not re.search(r"\byk\b.*\bucreti\b", normalized):
            continue
        candidates = [
            (item.center_x, amount)
            for item in line
            if (amount := _parse_money(item.text)) is not None
            and width * 0.28 <= item.center_x <= width * 0.50
        ]
        if not candidates:
            raise ValueError("YK Ücreti satırındaki tutar okunamadı.")
        yk_fee = max(candidates, key=lambda pair: pair[0])[1]
        break

    return base_gross + yk_fee, yk_fee


def _extract_additional_values(
    lines: list[list[TextBox]], width: float, height: float
) -> dict[str, Decimal]:
    gross_header_x = width * 0.40
    for line in lines:
        if (
            _line_y(line) < height * 0.25
            and "brut tutar" in _plain(_regional_text(line, width, 0.28, 0.48))
        ):
            matching = [item for item in line if "brut" in _plain(item.text)]
            if matching:
                gross_header_x = matching[0].center_x
            break

    overtime_y = _find_section_y(lines, "fazla mesailer", width, 0.00, 0.48)
    extra_y = _find_section_y(lines, "ek odemeler", width, 0.00, 0.48)
    gross_payments_y = _find_section_y(lines, "brut odemeler", width, 0.00, 0.48)
    if overtime_y is None or extra_y is None or gross_payments_y is None:
        raise ValueError("Fazla Mesailer veya Ek Ödemeler bölüm sınırları okunamadı.")

    overtime_total = _section_total(
        lines, width, overtime_y, extra_y, (0.00, 0.18), (0.28, 0.50), gross_header_x
    )
    extra_total = _section_total(
        lines, width, extra_y, gross_payments_y, (0.00, 0.18), (0.28, 0.50), gross_header_x
    )
    bonus = _named_amount(
        lines,
        "ikramiye",
        width,
        (0.00, 0.28),
        (0.28, 0.50),
        extra_y,
        gross_payments_y,
    )
    if overtime_total is None or extra_total is None:
        raise ValueError("Fazla Mesailer veya Ek Ödemeler toplamı okunamadı.")
    bonus = bonus if bonus is not None else Decimal("0")

    worker_x = _find_worker_x(lines, width)
    employer_x = _find_employer_x(lines, width)
    worker_sgk = _deduction_row_amount(lines, "sgk", worker_x, width, height)
    worker_sgdp = _deduction_row_amount(lines, "sgdp", worker_x, width, height)
    worker_unemployment = _deduction_row_amount(
        lines, "issizlik", worker_x, width, height
    )
    employer_sgk = _deduction_row_amount(lines, "sgk", employer_x, width, height)
    employer_sgdp = _deduction_row_amount(lines, "sgdp", employer_x, width, height)
    employer_unemployment = _deduction_row_amount(
        lines, "issizlik", employer_x, width, height
    )
    contribution_values = [
        worker_sgk,
        worker_sgdp,
        worker_unemployment,
        employer_sgk,
        employer_sgdp,
        employer_unemployment,
    ]
    if any(value is None for value in contribution_values):
        raise ValueError("İşçi veya işveren SGK/işsizlik tutarları okunamadı.")

    special_y = _find_section_y(lines, "ozel kesintiler", width, 0.43, 0.99)
    tax_discount_y = _find_section_y(lines, "vergi indirimi", width, 0.43, 0.99)
    incentives_y = _find_section_y(lines, "tesvikler", width, 0.43, 0.99)
    employer_cost_y = _find_section_y(
        lines, "isveren maliyeti", width, 0.43, 0.99
    )
    if None in (special_y, tax_discount_y, incentives_y, employer_cost_y):
        raise ValueError("Özel Kesintiler, Vergi İndirimi veya Teşvikler bölümü okunamadı.")

    special_total = _section_total(
        lines, width, special_y, tax_discount_y, (0.43, 0.67), (0.72, 0.99)
    )
    if special_total is None:
        raise ValueError("Özel Kesintiler bölüm toplamı okunamadı.")

    advance = Decimal("0")
    other_advances = Decimal("0")
    for line in lines:
        line_y = _line_y(line)
        if not special_y < line_y < tax_discount_y:
            continue
        label_text = _plain(_regional_text(line, width, 0.43, 0.70))
        if label_text != "avans" and not label_text.startswith("avans "):
            continue
        amounts = _regional_amounts(line, width, 0.70, 0.99)
        if not amounts:
            raise ValueError(f"{label_text.title()} satırındaki tutar okunamadı.")
        amount = max(amounts, key=lambda pair: pair[0])[1]
        if label_text == "avans":
            advance += amount
        else:
            other_advances += amount

    income_tax_incentive = _named_amount(
        lines,
        "gelir vergisi indirimi",
        width,
        (0.43, 0.78),
        (0.78, 0.99),
        tax_discount_y,
        incentives_y,
    )
    stamp_tax_incentive = _named_amount(
        lines,
        "damga vergisi indirimi",
        width,
        (0.43, 0.78),
        (0.78, 0.99),
        tax_discount_y,
        incentives_y,
    )
    net_paid = _named_amount(
        lines,
        "net odenen",
        width,
        (0.43, 0.76),
        (0.76, 0.99),
        employer_cost_y,
        height,
    )
    employer_cost = _named_amount(
        lines,
        "isveren maliyeti",
        width,
        (0.43, 0.78),
        (0.78, 0.99),
        incentives_y,
        height,
    )
    if None in (income_tax_incentive, stamp_tax_incentive, net_paid, employer_cost):
        raise ValueError("Vergi teşvikleri, net ödenen veya işveren maliyeti okunamadı.")

    incentives_total = Decimal("0")
    for line in lines:
        if not incentives_y < _line_y(line) < employer_cost_y:
            continue
        incentives_total += sum(
            (amount for _, amount in _regional_amounts(line, width, 0.78, 0.99)),
            Decimal("0"),
        )

    return {
        "overtime_total": overtime_total,
        "bonus": bonus,
        "other_extra_payments": extra_total - bonus,
        "worker_sgk": worker_sgk + worker_sgdp,
        "worker_unemployment": worker_unemployment,
        "employer_sgk": employer_sgk + employer_sgdp,
        "employer_unemployment": employer_unemployment,
        "income_tax_incentive": income_tax_incentive,
        "stamp_tax_incentive": stamp_tax_incentive,
        "advance": advance,
        "other_advances": other_advances,
        "account_369_deductions": special_total - advance - other_advances,
        "net_paid": net_paid,
        "incentives_total": incentives_total,
        "employer_cost": employer_cost,
    }


def extract_page_values(page) -> dict[str, Decimal]:
    native = _native_boxes(page)
    native_text = " ".join(item.text for item in native)
    if "gelir" in _plain(native_text) and "damga" in _plain(native_text):
        boxes = native
        width = float(page.rect.width)
        height = float(page.rect.height)
    else:
        boxes, width, height = _ocr_boxes(page)

    lines = _group_lines(boxes, height)
    gross, yk_fee = _extract_gross_amount(lines, width, height)
    values = _extract_additional_values(lines, width, height)
    worker_x = _find_worker_x(lines, width)
    income = _amount_on_named_line(lines, "gelir vergisi", worker_x, width)
    stamp = _amount_on_named_line(lines, "damga vergisi", worker_x, width)

    if income is None or stamp is None:
        fallback_income, fallback_stamp = _fallback_worker_amounts(lines, worker_x, width)
        income = income if income is not None else fallback_income
        stamp = stamp if stamp is not None else fallback_stamp

    missing = []
    if income is None:
        missing.append("gelir vergisi")
    if stamp is None:
        missing.append("damga vergisi")
    if missing:
        raise ValueError(f"İşçi sütunundaki {' ve '.join(missing)} okunamadı.")
    values.update({
        "gross": gross,
        "yk_fee": yk_fee,
        "income": income,
        "stamp": stamp,
    })
    return values


def analyze_tax_pdfs(pdf_files: list[tuple[str, str]]) -> dict:
    rows = []
    warnings = []
    totals = {key: Decimal("0") for key, _ in _EXCEL_COLUMNS}

    for path, original_name in pdf_files:
        try:
            document = fitz.open(path)
        except Exception as exc:
            warnings.append(f"{original_name}: PDF açılamadı ({exc}).")
            continue

        try:
            for page_number, page in enumerate(document, start=1):
                try:
                    values = extract_page_values(page)
                except RuntimeError:
                    raise
                except Exception as exc:
                    warnings.append(f"{original_name} - Sayfa {page_number}: {exc}")
                    continue
                for key in totals:
                    totals[key] += values[key]
                row = {
                    "filename": original_name,
                    "page": page_number,
                }
                row.update(values)
                rows.append(row)
        finally:
            document.close()

    result = {
        "rows": rows,
        "warnings": warnings,
        "totals": totals,
    }
    for key, value in totals.items():
        result[f"{key}_total"] = value
        result[f"{key}_total_text"] = format_tr_money(value)
    return result


def create_tax_excel(result: dict, output_dir: str) -> str:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Vergi Toplamları"
    sheet.sheet_view.showGridLines = False

    dark_blue = "1F4E78"
    medium_blue = "D9EAF7"
    light_blue = "EAF3F8"
    white = "FFFFFF"
    border_color = "B8C7D1"
    thin_border = Border(
        left=Side(style="thin", color=border_color),
        right=Side(style="thin", color=border_color),
        top=Side(style="thin", color=border_color),
        bottom=Side(style="thin", color=border_color),
    )

    last_column = 2 + len(_EXCEL_COLUMNS)
    last_column_letter = get_column_letter(last_column)
    sheet.merge_cells(f"A1:{last_column_letter}1")
    title = sheet["A1"]
    title.value = "Bordro Toplamları"
    title.fill = PatternFill("solid", fgColor=dark_blue)
    title.font = Font(name="Arial", size=15, bold=True, color=white)
    title.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 28

    header_row = 3
    headers = ["Dosya", "Sayfa"] + [header for _, header in _EXCEL_COLUMNS]
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(header_row, column, header)
        cell.fill = PatternFill("solid", fgColor=dark_blue)
        cell.font = Font(name="Arial", bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    first_data_row = header_row + 1
    for row_number, row in enumerate(result["rows"], start=first_data_row):
        values = [row["filename"], row["page"]] + [
            float(row[key]) for key, _ in _EXCEL_COLUMNS
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_number, column, value)
            cell.font = Font(name="Arial", size=10)
            cell.border = thin_border
            cell.alignment = Alignment(
                horizontal="right" if column >= 2 else "left",
                vertical="center",
            )
            if row_number % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=light_blue)
        for column in range(3, last_column + 1):
            sheet.cell(row_number, column).number_format = "#,##0.00"

    last_data_row = first_data_row + len(result["rows"]) - 1
    total_row = last_data_row + 1
    sheet.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=2)
    total_label = sheet.cell(total_row, 1, "GENEL TOPLAM")
    total_label.alignment = Alignment(horizontal="right", vertical="center")
    for column in range(3, last_column + 1):
        column_letter = get_column_letter(column)
        sheet.cell(
            total_row,
            column,
            f"=SUM({column_letter}{first_data_row}:{column_letter}{last_data_row})",
        )
    for column in range(1, last_column + 1):
        cell = sheet.cell(total_row, column)
        cell.fill = PatternFill("solid", fgColor=medium_blue)
        cell.font = Font(name="Arial", size=11, bold=True, color=dark_blue)
        cell.border = thin_border
    for column in range(3, last_column + 1):
        sheet.cell(total_row, column).number_format = "#,##0.00"

    sheet.auto_filter.ref = f"A{header_row}:{last_column_letter}{last_data_row}"
    sheet.freeze_panes = f"C{first_data_row}"
    sheet.row_dimensions[header_row].height = 42
    widths = [36, 9] + [22] * len(_EXCEL_COLUMNS)
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    for cell in sheet[header_row]:
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.print_title_rows = f"1:{header_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

    if result["warnings"]:
        warning_sheet = workbook.create_sheet("Uyarılar")
        warning_sheet.sheet_view.showGridLines = False
        warning_sheet["A1"] = "Kontrol Gereken Sayfalar"
        warning_sheet["A1"].fill = PatternFill("solid", fgColor="9C6500")
        warning_sheet["A1"].font = Font(name="Arial", size=13, bold=True, color=white)
        warning_sheet["A1"].alignment = Alignment(vertical="center")
        warning_sheet.row_dimensions[1].height = 25
        for row_number, warning in enumerate(result["warnings"], start=3):
            cell = warning_sheet.cell(row_number, 1, warning)
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        warning_sheet.column_dimensions["A"].width = 110
        warning_sheet.freeze_panes = "A3"

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"

    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(
        output_dir,
        f"bordro_toplamlari_{uuid.uuid4().hex[:8]}.xlsx",
    )
    workbook.save(output_path)
    return output_path
